"""실제 파일과 같은 레이아웃의 가상 엑셀(가상 이름·사번)을 만든다. 테스트와 사용법 시연용."""
from __future__ import annotations

from calendar import monthrange
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

YEAR = 2026
MONTH_COLS = [get_column_letter(c) for c in range(4, 16)]  # D..O

# (사번, 이름, 부서, 구분)
STAFF = [
    (900001, "가상사무1", "Finance", "OFFICE"),
    (900002, "가상사무2", "HR", "OFFICE"),
    (900003, "가상사무3", "Sales", "OFFICE"),
    (900004, "가상사무4", "Sales", "OFFICE"),
    (900101, "가상매장1", "Store A", "STORE"),
    (900102, "가상매장2", "Store B", "STORE"),
]
# 일자별 입력: {(사번): {(월, 일): 횟수}}
RIDES = {
    900001: {(7, 1): 2, (7, 2): 2, (7, 4): 2},  # 7/4는 토요일 → 제외
    900101: {(7, 6): 2, (7, 7): 2, (7, 12): 1},  # 7/6은 월요일(매장직 휴무) → 제외
    900002: {(8, 17): 2, (8, 18): 1, (8, 19): 3},  # 8/17 대체공휴일, 3은 잘못된 값
    900003: {(9, 1): 2, (9, 2): 2},
    900102: {(7, 14): 2, (8, 11): 2, (9, 8): 2},
}
EXTRA_NOT_IN_SUMMARY = (900999, "가상외부", "Sales", "OFFICE", {(7, 1): 2})


def build(path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Bike Commute Credits (sample)"
    ws["D10"] = YEAR
    for c, t in enumerate(["STAFF #", "NAME", "BUSINESS UNIT", "JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG",
                           "SEP", "OCT", "NOV", "DEC", "Q1", "Q2", "Q3", "Q4", "YTD"], start=1):
        ws.cell(11, c, t)
    rows = list(STAFF) + [STAFF[0]]  # 사번 중복 1건
    for i, (no, name, unit, kind) in enumerate(rows):
        r = 12 + i
        ws.cell(r, 1, no)
        ws.cell(r, 2, name)
        ws.cell(r, 3, unit)
        for m in range(1, 13):
            col = MONTH_COLS[m - 1]
            ws[f"{col}{r}"] = (f"=IFERROR(_xlfn.XLOOKUP(A{r},'{m:02d}_OFFICE'!A:A,'{m:02d}_OFFICE'!AJ:AJ),"
                               f"_xlfn.XLOOKUP(A{r},'{m:02d}_STORE'!A:A,'{m:02d}_STORE'!AJ:AJ))")
        for q, c in enumerate("PQRS"):
            a, b = MONTH_COLS[q * 3], MONTH_COLS[q * 3 + 2]
            ws[f"{c}{r}"] = f"=SUM({a}{r}:{b}{r})"
        ws[f"T{r}"] = f"=+SUM(P{r}:S{r})"

    for m in range(1, 13):
        dim = monthrange(YEAR, m)[1]
        for kind in ("OFFICE", "STORE"):
            sh = wb.create_sheet(f"{m:02d}_{kind}")
            sh["A1"] = "Bike Commute Credits (sample)"
            for c, t in enumerate(["STAFF #", "NAME", "BUSINESS UNIT", "others", YEAR], start=1):
                sh.cell(10, c, t)
            sh.cell(11, 5, m)
            for d in range(1, dim + 1):
                sh.cell(12, 4 + d, d)
            total_col = 5 + dim
            sh.cell(12, total_col, "Total")
            people = [s for s in STAFF if s[3] == kind]
            if kind == "OFFICE" and m == 7:
                people = people + [EXTRA_NOT_IN_SUMMARY[:4]]
            for i, (no, name, unit, _) in enumerate(people):
                r = 13 + i
                sh.cell(r, 1, no)
                sh.cell(r, 2, name)
                sh.cell(r, 3, unit)
                data = RIDES.get(no, {})
                if no == EXTRA_NOT_IN_SUMMARY[0]:
                    data = EXTRA_NOT_IN_SUMMARY[4]
                for (mm, dd), v in data.items():
                    if mm == m:
                        sh.cell(r, 4 + dd, v)
                last_col = get_column_letter(4 + dim)
                sh.cell(r, total_col, f"=+SUM(E{r}:{last_col}{r})")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys

    out = build(sys.argv[1] if len(sys.argv) > 1 else "sample/sample.xlsx")
    print(out)
