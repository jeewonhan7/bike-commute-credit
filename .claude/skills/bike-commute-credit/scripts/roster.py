"""입사자/퇴사자 명단 반영: HR 명단 CSV와 엑셀을 비교하고, 신규 입사자 행을 추가한 사본을 만든다."""
from __future__ import annotations

import re
from copy import copy
from pathlib import Path

import openpyxl

import cclib
from cclib import Book, RosterRow, _find_layout, clean, norm_staff

MONTH_COLS = 12  # Summary D..O


def diff(book: Book, roster: dict[str, RosterRow]) -> dict:
    new_hires = [r for no, r in roster.items() if no not in book.master and not r.leave]
    leavers = [r for no, r in roster.items() if r.leave and no in book.master]
    name_changes = [(book.master[no], r) for no, r in roster.items()
                    if no in book.master and r.name and book.master[no].name != r.name]
    unit_changes = [(book.master[no], r) for no, r in roster.items()
                    if no in book.master and r.unit and book.master[no].unit != r.unit]
    not_in_roster = [p for no, p in book.master.items() if no not in roster]
    return {"new_hires": new_hires, "leavers": leavers, "name_changes": name_changes,
            "unit_changes": unit_changes, "not_in_roster": not_in_roster}


def write_diff(path, d: dict) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "신규 입사자"
    ws.append(["사번", "이름", "부서", "구분", "입사일"])
    for r in d["new_hires"]:
        ws.append([r.staff_no, r.name, r.unit, r.kind, r.join.isoformat() if r.join else ""])
    ws = wb.create_sheet("퇴사자")
    ws.append(["사번", "이름", "부서", "퇴사일"])
    for r in d["leavers"]:
        ws.append([r.staff_no, r.name, r.unit, r.leave.isoformat()])
    ws = wb.create_sheet("이름·부서 변경")
    ws.append(["사번", "엑셀 이름", "명단 이름", "엑셀 부서", "명단 부서"])
    seen = set()
    for p, r in d["name_changes"] + d["unit_changes"]:
        if p.staff_no not in seen:
            seen.add(p.staff_no)
            ws.append([p.staff_no, p.name, r.name, p.unit, r.unit])
    ws = wb.create_sheet("명단에 없는 직원")
    ws.append(["사번", "이름", "부서", "Summary 행"])
    for p in d["not_in_roster"]:
        ws.append([p.staff_no, p.name, p.unit, p.row])
    for s in wb.worksheets:
        for col in "ABCDE":
            s.column_dimensions[col].width = 16
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _swap_row(formula: str, old: int, new: int) -> str:
    return re.sub(rf"([A-Z]{{1,3}}){old}(?!\d)", rf"\g<1>{new}", formula)


def _copy_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)


def apply_new_hires(book: Book, new_hires: list[RosterRow], out_path) -> list[str]:
    """신규 입사자를 Summary와 입사월 이후의 월별 시트에 추가한 사본을 저장한다. 원본은 건드리지 않는다."""
    wb = openpyxl.load_workbook(book.path)
    ws = wb["Summary"]
    tmpl = next((p.row for p in book.master.values()
                 if isinstance(ws.cell(p.row, 4).value, str) and ws.cell(p.row, 4).value.startswith("=")), None)
    if tmpl is None:
        raise SystemExit("Summary에서 수식이 들어 있는 직원 행을 찾지 못해 신규 행의 서식을 만들 수 없습니다.")
    log = []
    next_row = max(p.row for p in book.master.values()) + 1
    for h in new_hires:
        if h.join and h.join.year > book.year:
            log.append(f"{h.staff_no} {h.name}: 입사일({h.join})이 {book.year}년 이후라 건너뜀")
            continue
        first_month = 1 if (not h.join or h.join.year < book.year) else h.join.month
        r = next_row
        next_row += 1
        for c in range(1, ws.max_column + 1):
            _copy_style(ws.cell(tmpl, c), ws.cell(r, c))
        ws.cell(r, 1, int(h.staff_no))
        ws.cell(r, 2, h.name)
        ws.cell(r, 3, h.unit)
        for m in range(1, MONTH_COLS + 1):
            col = 3 + m
            ws.cell(r, col, 0 if m < first_month else _swap_row(ws.cell(tmpl, col).value, tmpl, r))
        for col in range(16, 21):  # Q1..Q4, YTD
            v = ws.cell(tmpl, col).value
            if isinstance(v, str) and v.startswith("="):
                ws.cell(r, col, _swap_row(v, tmpl, r))
        log.append(f"{h.staff_no} {h.name}: Summary {r}행 추가")

        for m in range(first_month, 13):
            name = cclib.sheet_name(m, h.kind)
            if name not in wb.sheetnames:
                continue
            sh = wb[name]
            layout = _find_layout(sh)
            if layout is None:
                log.append(f"{name}: 일자 헤더를 못 찾아 건너뜀")
                continue
            hdr, _, total_col = layout
            last = max((x for x in range(hdr + 1, sh.max_row + 1) if norm_staff(sh.cell(x, 1).value)), default=hdr)
            new = last + 1
            for c in range(1, sh.max_column + 1):
                _copy_style(sh.cell(last, c), sh.cell(new, c))
            sh.cell(new, 1, int(h.staff_no))
            sh.cell(new, 2, h.name)
            sh.cell(new, 3, h.unit)
            if total_col and isinstance(sh.cell(last, total_col).value, str):
                sh.cell(new, total_col, _swap_row(sh.cell(last, total_col).value, last, new))
        log.append(f"{h.staff_no} {h.name}: {first_month}월~12월 {h.kind} 시트 추가")
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return log
