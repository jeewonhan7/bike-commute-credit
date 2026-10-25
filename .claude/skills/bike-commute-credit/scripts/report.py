"""정산 엑셀과 검증 오류 목록 엑셀 출력."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import cclib
from cclib import QUARTER_MONTHS, SEV_ORDER, deadline, payout_date

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEV_FILL = {"ERROR": PatternFill("solid", fgColor="F8CBAD"), "WARN": PatternFill("solid", fgColor="FFE699")}
MONEY = "#,##0"


def _header(ws, row, titles):
    for c, t in enumerate(titles, start=1):
        cell = ws.cell(row, c, t)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_settlement(path, book, quarter, rows, as_of: date, voids_count: int) -> None:
    months = QUARTER_MONTHS[quarter]
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "정산"
    titles = (["사번", "이름", "부서", "구분"]
              + [f"{m}월 입력" for m in months] + [f"{m}월 인정" for m in months]
              + ["인정 횟수", "산정액", "기지급액", "이번 지급액", "누적 지급액", "잔여 한도", "비고"])
    _header(ws, 1, titles)
    rows = sorted(rows, key=lambda r: (cclib.kind_order(r["kind"]), r["unit"], r["name"]))
    for i, r in enumerate(rows, start=2):
        vals = ([r["staff_no"], r["name"], r["unit"], r["kind"]]
                + [r["raw"][m] for m in months] + [r["credited"][m] for m in months]
                + [r["rides"], r["gross"], r["paid_before"], r["payout"], r["cum_after"], r["remaining"], r["note"]])
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(i, c, v)
            if isinstance(v, int) and c > 10:
                cell.number_format = MONEY
        if r["voided"]:
            for c in range(1, len(vals) + 1):
                ws.cell(i, c).fill = SEV_FILL["WARN"]
    last = len(rows) + 1
    total_row = last + 1
    ws.cell(total_row, 1, "합계").font = Font(bold=True)
    pay_col = 4 + 2 * len(months) + 4
    for c in list(range(5, 5 + 2 * len(months) + 1)) + [pay_col - 2, pay_col - 1, pay_col]:
        col = get_column_letter(c)
        cell = ws.cell(total_row, c, f"=SUM({col}2:{col}{last})")
        cell.font = Font(bold=True)
        cell.number_format = MONEY
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{last}"
    _widths(ws, [10, 10, 18, 7] + [9] * (2 * len(months)) + [9, 12, 12, 13, 13, 12, 50])

    by_unit = defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        u = by_unit[(r["kind"], r["unit"])]
        u[0] += 1
        u[1] += 1 if r["payout"] > 0 else 0
        u[2] += r["rides"]
        u[3] += r["payout"]
    ws2 = wb.create_sheet("부서별 요약")
    _header(ws2, 1, ["구분", "부서", "인원", "지급 인원", "인정 횟수", "지급액"])
    for i, ((kind, unit), v) in enumerate(sorted(by_unit.items(), key=lambda kv: (cclib.kind_order(kv[0][0]), kv[0][1])), start=2):
        for c, x in enumerate([kind, unit] + v, start=1):
            cell = ws2.cell(i, c, x)
            if c == 6:
                cell.number_format = MONEY
    n = len(by_unit) + 1
    ws2.cell(n + 1, 1, "합계").font = Font(bold=True)
    for c in range(3, 7):
        col = get_column_letter(c)
        cell = ws2.cell(n + 1, c, f"=SUM({col}2:{col}{n})")
        cell.font = Font(bold=True)
        cell.number_format = MONEY
    _widths(ws2, [8, 20, 8, 10, 10, 14])

    errs = Counter(i.severity for i in book.issues if i.month is None or i.month in months)
    ws3 = wb.create_sheet("기준")
    info = [
        ("대상", f"{book.year}년 {quarter}분기 ({months[0]}~{months[-1]}월)"),
        ("크레딧 지급일", payout_date(book.year, quarter).isoformat()),
        ("입력 마감일", deadline(book.year, quarter).isoformat()),
        ("정산 기준일", as_of.isoformat()),
        ("1회당 크레딧", cclib.RATE),
        ("연 한도", cclib.ANNUAL_CAP),
        ("원본 파일", Path(book.path).name),
        ("검증 ERROR", errs.get("ERROR", 0)),
        ("검증 WARN", errs.get("WARN", 0)),
        ("스크리닝 불일치 건수(무효 월)", voids_count),
        ("지급 대상 인원", sum(1 for r in rows if r["payout"] > 0)),
        ("총 지급액", sum(r["payout"] for r in rows)),
        ("주의", "금액 지급 전 검증 오류 목록과 스크리닝 결과를 HR이 최종 확인해야 합니다."),
    ]
    for i, (k, v) in enumerate(info, start=1):
        ws3.cell(i, 1, k).font = Font(bold=True)
        cell = ws3.cell(i, 2, v)
        if isinstance(v, int) and k in ("1회당 크레딧", "연 한도", "총 지급액"):
            cell.number_format = MONEY
    _widths(ws3, [30, 70])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_errors(path, book, months=None) -> int:
    issues = [i for i in book.issues if months is None or i.month is None or i.month in months]
    issues.sort(key=lambda i: (SEV_ORDER[i.severity], i.code, i.sheet, i.staff_no, i.date or date.min))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "검증 오류 목록"
    _header(ws, 1, ["심각도", "코드", "사번", "이름", "부서", "시트", "셀", "날짜", "내용", "처리"])
    for i, it in enumerate(issues, start=2):
        vals = [it.severity, it.code, it.staff_no, it.name, it.unit, it.sheet, it.cell,
                it.date.isoformat() if it.date else "", it.message, it.action]
        for c, v in enumerate(vals, start=1):
            ws.cell(i, c, v)
        if it.severity in SEV_FILL:
            ws.cell(i, 1).fill = SEV_FILL[it.severity]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(issues) + 1}"
    _widths(ws, [9, 18, 10, 10, 16, 10, 8, 12, 80, 12])

    ws2 = wb.create_sheet("요약")
    _header(ws2, 1, ["심각도", "코드", "건수"])
    cnt = Counter((i.severity, i.code) for i in issues)
    for r, ((sev, code), n) in enumerate(sorted(cnt.items(), key=lambda kv: (SEV_ORDER[kv[0][0]], kv[0][1])), start=2):
        for c, v in enumerate([sev, code, n], start=1):
            ws2.cell(r, c, v)
    _widths(ws2, [9, 20, 8])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(issues)
