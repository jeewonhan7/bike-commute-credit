"""자전거 통근 크레딧 정산 핵심 로직: 엑셀 파싱, 규칙 검증, 정산 계산.

엑셀 구조:
  - Summary: 11행 헤더(STAFF # / NAME / BUSINESS UNIT / JAN..DEC / Q1..Q4 / YTD), 12행부터 직원 마스터
  - MM_<시트 접미사>(예: 01_OFFICE, 01_STORE): 헤더 행에 E열부터 1..31일, 마지막에 Total
    각 일자 칸 = 그날의 라이딩 횟수(왕복 2, 편도 1)

규칙(단가, 연 한도, 근무요일, 마감일, 지급일)은 references/config.json 에서 읽는다.
"""
from __future__ import annotations

import csv
import json
import re
from calendar import monthrange
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

DEFAULT_CONFIG = Path(__file__).resolve().parent.parent / "references" / "config.json"
QUARTER_MONTHS = {1: (1, 2, 3), 2: (4, 5, 6), 3: (7, 8, 9), 4: (10, 11, 12)}
FIRST_DAY_COL = 5  # E열
WEEKDAY_KO = "월화수목금토일"
SEV_ORDER = {"ERROR": 0, "WARN": 1, "INFO": 2}

# set_config()가 채운다
CFG: dict = {}
RATE = 0
ANNUAL_CAP = 0
KINDS: dict = {}
SHEET_RE = re.compile(r"$^")
_SUFFIX_TO_KIND: dict = {}


def set_config(path=None) -> dict:
    """설정 JSON을 읽어 모듈 전역에 반영한다. 경로가 없으면 references/config.json."""
    global CFG, RATE, ANNUAL_CAP, KINDS, SHEET_RE, _SUFFIX_TO_KIND
    path = Path(path) if path else DEFAULT_CONFIG
    try:
        cfg = json.loads(path.read_text(encoding="utf-8-sig"))
        kinds = cfg["kinds"]
        rate, cap = int(cfg["rate_per_ride"]), int(cfg["annual_cap"])
        int(cfg["deadline_day"]), int(cfg["payout_day"])
        for k, v in kinds.items():
            v["sheet_suffix"], v["workdays"]
    except (OSError, KeyError, ValueError, TypeError) as e:
        raise SystemExit(f"설정 파일을 읽을 수 없습니다({path}): {e!r}. references/config.json 형식을 확인하세요.")
    CFG, KINDS, RATE, ANNUAL_CAP = cfg, kinds, rate, cap
    _SUFFIX_TO_KIND = {v["sheet_suffix"]: k for k, v in kinds.items()}
    SHEET_RE = re.compile(r"^(\d{2})_(" + "|".join(re.escape(s) for s in _SUFFIX_TO_KIND) + r")$")
    return cfg


def kind_order(kind: str) -> int:
    keys = list(KINDS)
    return keys.index(kind) if kind in keys else len(keys)


def kind_label(kind: str) -> str:
    return KINDS.get(kind, {}).get("label", kind)


def sheet_name(month: int, kind: str) -> str:
    return f"{month:02d}_{KINDS[kind]['sheet_suffix']}"


set_config()


def quarter_of(month: int) -> int:
    return (month - 1) // 3 + 1


def deadline(year: int, q: int) -> date:
    """분기 마지막 달의 익월 deadline_day일."""
    last = QUARTER_MONTHS[q][-1]
    d = int(CFG["deadline_day"])
    return date(year + 1, 1, d) if last == 12 else date(year, last + 1, d)


def payout_date(year: int, q: int) -> date:
    """크레딧 지급일: 분기 마지막 달의 익월 payout_day일(Q4는 익년 1월)."""
    last = QUARTER_MONTHS[q][-1]
    d = int(CFG["payout_day"])
    return date(year + 1, 1, d) if last == 12 else date(year, last + 1, d)


def norm_staff(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\D", "", str(v))
    return s or None


def clean(v) -> str:
    return "" if v is None else str(v).strip()


def parse_date(s: str) -> date | None:
    s = clean(s)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"날짜 형식을 읽을 수 없습니다: {s!r} (YYYY-MM-DD 권장)")


# --------------------------------------------------------------------------- 모델
@dataclass
class Person:
    staff_no: str
    name: str
    unit: str
    row: int


@dataclass
class Entry:
    staff_no: str
    name: str
    unit: str
    kind: str
    date: date
    rides: int
    sheet: str
    row: int
    cell: str
    valid: bool = True
    reason: str = ""


@dataclass
class Issue:
    severity: str
    code: str
    message: str
    action: str = "확인필요"
    staff_no: str = ""
    name: str = ""
    unit: str = ""
    sheet: str = ""
    cell: str = ""
    date: date | None = None
    month: int | None = None


@dataclass
class RosterRow:
    staff_no: str
    name: str
    unit: str
    kind: str
    join: date | None
    leave: date | None


@dataclass
class Book:
    path: str
    year: int
    master: dict = field(default_factory=dict)  # staff_no -> Person
    rows: list = field(default_factory=list)  # 월별 시트의 직원 행
    entries: list = field(default_factory=list)
    issues: list = field(default_factory=list)
    sheets: dict = field(default_factory=dict)  # sheet -> (kind, month)

    def add(self, severity, code, message, action="확인필요", **kw):
        self.issues.append(Issue(severity, code, message, action, **kw))


# --------------------------------------------------------------------------- 로딩
def _find_layout(ws):
    """일자 헤더 행(E=1, F=2)과 일자별 열, Total 열을 찾는다."""
    hdr = None
    for r in range(1, min(ws.max_row, 30) + 1):
        if ws.cell(r, FIRST_DAY_COL).value == 1 and ws.cell(r, FIRST_DAY_COL + 1).value == 2:
            hdr = r
            break
    if hdr is None:
        return None
    day_cols, total_col = {}, None
    for c in range(FIRST_DAY_COL, ws.max_column + 1):
        v = ws.cell(hdr, c).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            day_cols[int(v)] = c
        elif isinstance(v, str) and v.strip().lower() == "total":
            total_col = c
            break
    return hdr, day_cols, total_col


def load_book(path, year: int | None = None) -> Book:
    wb = openpyxl.load_workbook(path)
    if "Summary" not in wb.sheetnames:
        raise SystemExit("Summary 시트가 없습니다. 정산 엑셀 형식이 맞는지 확인하세요.")
    ws = wb["Summary"]
    if year is None:
        v = ws["D10"].value
        if not isinstance(v, int):
            raise SystemExit("연도를 찾을 수 없습니다. --year 로 지정하세요.")
        year = v
    book = Book(path=str(path), year=year)

    hdr = next((r for r in range(1, 40) if clean(ws.cell(r, 1).value) == "STAFF #"), None)
    if hdr is None:
        raise SystemExit("Summary 시트에서 'STAFF #' 헤더를 찾을 수 없습니다.")
    for r in range(hdr + 1, ws.max_row + 1):
        no = norm_staff(ws.cell(r, 1).value)
        if not no:
            continue
        p = Person(no, clean(ws.cell(r, 2).value), clean(ws.cell(r, 3).value), r)
        if no in book.master:
            first = book.master[no]
            book.add("WARN", "MASTER_DUP",
                     f"Summary에 같은 사번이 {first.row}행과 {r}행에 중복됩니다. 첫 행 기준으로 1명으로 계산합니다.",
                     "중복 행 정리", staff_no=no, name=p.name, unit=p.unit, sheet="Summary", cell=f"A{r}")
            continue
        book.master[no] = p

    month_sheets = sorted(
        (n for n in wb.sheetnames if SHEET_RE.match(n)),
        key=lambda n: (int(n[:2]), kind_order(_SUFFIX_TO_KIND[n[3:]])),
    )
    seen: dict[tuple[int, str], tuple[str, int]] = {}
    for name in month_sheets:
        m = SHEET_RE.match(name)
        month, kind = int(m.group(1)), _SUFFIX_TO_KIND[m.group(2)]
        book.sheets[name] = (kind, month)
        sh = wb[name]
        layout = _find_layout(sh)
        if layout is None:
            book.add("ERROR", "SHEET_LAYOUT", "일자 헤더(1, 2, 3 ...)를 찾을 수 없어 시트를 읽지 못했습니다.",
                     "시트 형식 확인", sheet=name, month=month)
            continue
        hdr_row, day_cols, _ = layout
        dim = monthrange(book.year, month)[1]
        if set(day_cols) != set(range(1, dim + 1)):
            book.add("WARN", "SHEET_LAYOUT",
                     f"{month}월은 {dim}일까지인데 시트의 일자 열은 {min(day_cols)}~{max(day_cols)}일입니다.",
                     "시트 형식 확인", sheet=name, month=month)
        for r in range(hdr_row + 1, sh.max_row + 1):
            no = norm_staff(sh.cell(r, 1).value)
            if not no:
                continue
            nm, unit = clean(sh.cell(r, 2).value), clean(sh.cell(r, 3).value)
            row_info = {"sheet": name, "row": r, "staff_no": no, "name": nm, "unit": unit,
                        "month": month, "kind": kind, "dup": False}
            first = seen.setdefault((month, no), (name, r))
            if first != (name, r):
                row_info["dup"] = True
                book.add("ERROR", "DUP_IN_MONTH",
                         f"같은 달에 사번이 중복 입력되었습니다(최초 입력: {first[0]} {first[1]}행). 이 행은 계산에서 제외합니다.",
                         "제외", staff_no=no, name=nm, unit=unit, sheet=name, cell=f"A{r}", month=month)
            book.rows.append(row_info)
            for day, col in day_cols.items():
                if day > dim:
                    continue
                raw = sh.cell(r, col).value
                if raw is None or (isinstance(raw, str) and not raw.strip()):
                    continue
                coord = sh.cell(r, col).coordinate
                d = date(book.year, month, day)
                rides, ok = _ride_value(raw, book, no, nm, unit, name, coord, d)
                if rides == 0:
                    continue
                e = Entry(no, nm, unit, kind, d, rides, name, r, coord, valid=ok,
                          reason="" if ok else "값 오류")
                if row_info["dup"]:
                    e.valid, e.reason = False, "DUP_IN_MONTH"
                book.entries.append(e)
    return book


def _ride_value(raw, book, no, nm, unit, sheet, coord, d):
    """칸 값을 라이딩 횟수로 변환. (횟수, 유효여부)"""
    kw = dict(staff_no=no, name=nm, unit=unit, sheet=sheet, cell=coord, date=d, month=d.month)
    val = raw
    if isinstance(raw, str):
        try:
            val = float(raw.strip())
            book.add("WARN", "TEXT_NUMBER", f"숫자가 텍스트로 입력되어 있습니다({raw!r}). 숫자로 인정합니다.", "형식 정리", **kw)
        except ValueError:
            book.add("ERROR", "INVALID_VALUE", f"숫자가 아닌 값이 입력되어 있습니다({raw!r}).", "제외", **kw)
            return 1, False
    if isinstance(val, bool) or not isinstance(val, (int, float)):
        book.add("ERROR", "INVALID_VALUE", f"숫자가 아닌 값이 입력되어 있습니다({raw!r}).", "제외", **kw)
        return 1, False
    if val == 0:
        return 0, True
    if val not in (1, 2):
        book.add("ERROR", "INVALID_VALUE", f"라이딩 횟수는 1(편도) 또는 2(왕복)만 가능합니다(입력값 {raw}).", "제외", **kw)
        return int(val) if float(val).is_integer() else 1, False
    return int(val), True


def load_holidays(path, year: int) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    items = data.get(str(year))
    if items is None:
        raise SystemExit(f"{path} 에 {year}년 공휴일 목록이 없습니다. 파일에 추가하거나 --holidays 로 다른 파일을 지정하세요.")
    return {date.fromisoformat(i["date"]): i["name"] for i in items}


def infer_kind(unit: str) -> str:
    """명단에 type이 비어 있을 때: unit_suffixes가 맞는 구분, 없으면 첫 번째 구분."""
    for k, v in KINDS.items():
        if any(unit.endswith(s) for s in v.get("unit_suffixes", [])):
            return k
    return next(iter(KINDS))


def load_roster(path) -> dict[str, RosterRow]:
    """명단 CSV: staff_no,name,unit,type,join_date,leave_date  (type = config의 구분 이름 또는 시트 접미사)"""
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.DictReader(f), start=2):
            no = norm_staff(row.get("staff_no"))
            if not no:
                continue
            unit = clean(row.get("unit"))
            kind = clean(row.get("type")).upper()
            kind = kind if kind in KINDS else _SUFFIX_TO_KIND.get(kind) or infer_kind(unit)
            try:
                join, leave = parse_date(row.get("join_date", "")), parse_date(row.get("leave_date", ""))
            except ValueError as ex:
                raise SystemExit(f"{path} {i}행: {ex}")
            out[no] = RosterRow(no, clean(row.get("name")), unit, kind, join, leave)
    return out


# --------------------------------------------------------------------------- 검증
def workday_problem(kind: str, d: date, holidays: dict) -> str | None:
    """근무일이 아니면 사유 문자열, 근무일이면 None. 근무 요일과 공휴일 휴무 여부는 config."""
    k = KINDS[kind]
    label, wd = k.get("label", kind), WEEKDAY_KO[d.weekday()]
    if d.weekday() not in k["workdays"]:
        days = "".join(WEEKDAY_KO[i] for i in k["workdays"])
        return f"{label}({kind})은 {days}요일 근무인데 {wd}요일입니다."
    if k.get("off_on_holidays") and d in holidays:
        return f"{label}({kind})은 공휴일({holidays[d]})에 근무하지 않습니다."
    return None


def validate(book: Book, holidays: dict, as_of: date, roster: dict | None = None,
             snapshot: dict | None = None) -> None:
    """엔트리별 규칙 검증. 위반 엔트리는 valid=False 로 표시하고 이슈를 쌓는다."""
    rows_with_rides = {(e.sheet, e.row) for e in book.entries}
    orphans = defaultdict(list)
    for r in book.rows:
        if r["staff_no"] not in book.master:
            orphans[r["staff_no"]].append(r)
            continue
        p = book.master.get(r["staff_no"])
        if p and p.name and r["name"] and p.name != r["name"]:
            book.add("WARN", "NAME_MISMATCH", f"Summary 이름({p.name})과 시트 이름({r['name']})이 다릅니다.", "확인",
                     staff_no=r["staff_no"], name=r["name"], unit=r["unit"], sheet=r["sheet"],
                     cell=f"B{r['row']}", month=r["month"])
        if p and p.unit and r["unit"] and p.unit != r["unit"]:
            book.add("WARN", "UNIT_MISMATCH", f"Summary 부서({p.unit})와 시트 부서({r['unit']})가 다릅니다.", "확인",
                     staff_no=r["staff_no"], name=r["name"], unit=r["unit"], sheet=r["sheet"],
                     cell=f"C{r['row']}", month=r["month"])
    for no, rs in orphans.items():
        with_rides = [r for r in rs if (r["sheet"], r["row"]) in rows_with_rides]
        first = (with_rides or rs)[0]
        if with_rides:
            book.add("ERROR", "NOT_IN_SUMMARY",
                     f"Summary 명단에 없는 사번인데 월별 시트({len(with_rides)}개 시트)에 라이딩이 입력되어 있습니다. 정산에서 제외합니다.",
                     "제외", staff_no=no, name=first["name"], unit=first["unit"], sheet=first["sheet"],
                     cell=f"A{first['row']}", month=first["month"] if len(with_rides) == 1 else None)
        else:
            book.add("WARN", "NOT_IN_SUMMARY",
                     f"Summary 명단에 없는 사번이 월별 시트({len(rs)}개 시트)에 남아 있습니다(라이딩 입력 없음). 퇴사자 정리 대상인지 확인하세요.",
                     "확인", staff_no=no, name=first["name"], unit=first["unit"], sheet=first["sheet"],
                     cell=f"A{first['row']}")
    kinds = defaultdict(set)
    for r in book.rows:
        kinds[r["staff_no"]].add(r["kind"])
    for no, ks in kinds.items():
        if len(ks) > 1 and no in book.master:
            p = book.master[no]
            book.add("WARN", "KIND_CONFLICT", "여러 구분(예: 사무직과 매장직) 시트에 모두 나타납니다. 근무일 규칙이 시트 기준으로 적용됩니다.",
                     "확인", staff_no=no, name=p.name, unit=p.unit)

    for e in book.entries:
        if not e.valid:
            continue
        kw = dict(staff_no=e.staff_no, name=e.name, unit=e.unit, sheet=e.sheet, cell=e.cell, date=e.date,
                  month=e.date.month)
        if e.staff_no not in book.master:
            e.valid, e.reason = False, "NOT_IN_SUMMARY"
            continue
        if e.date > as_of:
            _bad(book, e, "FUTURE_DATE", f"기준일({as_of})보다 미래 날짜에 라이딩이 입력되었습니다.", kw)
            continue
        why = workday_problem(e.kind, e.date, holidays)
        if why:
            _bad(book, e, "NON_WORKDAY", f"근무일이 아닌 날의 라이딩은 인정되지 않습니다. {why}", kw)
            continue
        if roster and e.staff_no in roster:
            rr = roster[e.staff_no]
            if rr.join and e.date < rr.join:
                _bad(book, e, "BEFORE_JOIN", f"입사일({rr.join}) 이전 날짜입니다.", kw)
                continue
            if rr.leave and e.date > rr.leave:
                _bad(book, e, "AFTER_LEAVE", f"퇴사일({rr.leave}) 이후 날짜입니다.", kw)
                continue
        if snapshot is not None and snapshot.get((e.staff_no, e.date), 0) < e.rides:
            _bad(book, e, "LATE_ENTRY", "마감일 스냅샷에 없는(또는 횟수가 늘어난) 입력입니다. 마감 후 입력은 인정되지 않습니다.", kw)
            continue


def _bad(book, e, code, msg, kw):
    e.valid, e.reason = False, code
    book.add("ERROR", code, msg, "제외", **kw)


def snapshot_index(book: Book) -> dict:
    out = defaultdict(int)
    for e in book.entries:
        if e.valid:
            out[(e.staff_no, e.date)] += e.rides
    return dict(out)


def quarter_notes(book: Book, quarter: int, as_of: date, has_snapshot: bool) -> None:
    dl = deadline(book.year, quarter)
    if as_of <= dl:
        book.add("WARN", "QUARTER_OPEN",
                 f"{quarter}분기 입력 마감일({dl})이 아직 지나지 않았습니다(기준일 {as_of}). 마감 후 다시 정산하세요.", "마감 후 재실행")
    elif not has_snapshot:
        book.add("INFO", "LATE_UNVERIFIABLE",
                 f"엑셀에는 입력 시각이 없어 마감일({dl}) 이후 입력/수정 여부를 확인할 수 없습니다. "
                 "마감일에 저장해 둔 사본이 있으면 --snapshot 으로 지정하세요.", "참고")
    for m in QUARTER_MONTHS[quarter]:
        for kind in KINDS:
            if sheet_name(m, kind) not in book.sheets:
                book.add("WARN", "MISSING_SHEET", f"{sheet_name(m, kind)} 시트가 없습니다.", "확인", month=m)


# --------------------------------------------------------------------------- 정산
def load_screening(path) -> dict:
    """랜덤 스크리닝 결과 CSV: staff_no,month,result(PASS/FAIL/REVIEW),note → {(staff_no, month): row}"""
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            no = norm_staff(row.get("staff_no"))
            if not no:
                continue
            res = clean(row.get("result")).upper()
            if res in ("FAIL", "불일치", "상이", "무효"):
                res = "FAIL"
            elif res in ("REVIEW", "확인필요"):
                res = "REVIEW"
            else:
                res = "PASS"
            out[(no, int(row["month"]))] = {"result": res, "note": clean(row.get("note"))}
    return out


def load_prior_paid(path) -> dict:
    """이전 분기까지 실제 지급액 CSV: staff_no,paid"""
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            no = norm_staff(row.get("staff_no"))
            if no:
                out[no] = int(float(row["paid"]))
    return out


def settle(book: Book, quarter: int, screening: dict | None = None, prior_paid: dict | None = None) -> list[dict]:
    """분기 정산. 연 한도까지 누적 지급하고, 스크리닝 불일치 월은 전체 무효."""
    screening = screening or {}
    voided = {k for k, v in screening.items() if v["result"] == "FAIL"}
    raw, valid = defaultdict(int), defaultdict(int)
    kind_of = {}
    for r in book.rows:
        kind_of.setdefault(r["staff_no"], r["kind"])
    for e in book.entries:
        raw[(e.staff_no, e.date.month)] += e.rides
        if e.valid:
            valid[(e.staff_no, e.date.month)] += e.rides
    err_count = defaultdict(int)
    for i in book.issues:
        if i.severity == "ERROR" and i.staff_no and (i.month is None or quarter_of(i.month) == quarter):
            err_count[i.staff_no] += 1

    rows = []
    for p in book.master.values():
        cum = 0
        for q in range(1, quarter + 1):
            months = QUARTER_MONTHS[q]
            if q == quarter and prior_paid and p.staff_no in prior_paid:
                cum = min(ANNUAL_CAP, prior_paid[p.staff_no])
            credited = {m: (0 if (p.staff_no, m) in voided else valid.get((p.staff_no, m), 0)) for m in months}
            gross = sum(credited.values()) * RATE
            pay = min(gross, max(0, ANNUAL_CAP - cum))
            before = cum
            cum += pay
        months = QUARTER_MONTHS[quarter]
        notes = []
        void_m = [m for m in months if (p.staff_no, m) in voided]
        if void_m:
            notes.append("스크리닝 불일치로 무효: " + ", ".join(f"{m}월" for m in void_m))
        if gross > pay:
            notes.append(f"연 한도 도달로 {gross - pay:,}원 미지급")
        elif cum >= ANNUAL_CAP:
            notes.append("연 한도 도달")
        if err_count[p.staff_no]:
            notes.append(f"검증 오류 {err_count[p.staff_no]}건(오류 목록 참고)")
        rows.append({
            "staff_no": p.staff_no, "name": p.name, "unit": p.unit, "kind": kind_of.get(p.staff_no, ""),
            "raw": {m: raw.get((p.staff_no, m), 0) for m in months},
            "credited": credited, "voided": void_m,
            "rides": sum(credited.values()), "gross": gross, "paid_before": before,
            "payout": pay, "cum_after": cum, "remaining": ANNUAL_CAP - cum, "note": "; ".join(notes),
        })
    return rows
